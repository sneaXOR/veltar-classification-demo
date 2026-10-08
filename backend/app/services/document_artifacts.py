from __future__ import annotations

import copy
import hashlib
import io
import os
import shutil
import tempfile
import textwrap
import time
from dataclasses import dataclass
from email import policy
from email.parser import BytesParser
from pathlib import Path
from typing import Any

import fitz
from docx import Document as DocxDocument
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Pt, RGBColor
from openpyxl import load_workbook
from PIL import Image, ImageDraw, ImageFont
from pptx import Presentation

from app.models.artifact import DocumentArtifactRevision
from app.models.document import Document, DocumentPage
from app.models.enums import ArtifactRevisionKind, ArtifactStatus, DocumentKind, DocumentStatus, JobStatus
from app.schemas.api import (
    DocumentArtifactResponse,
    DocumentAuditResponse,
    DocumentDecisionResponse,
    DocumentListItem,
    DocumentProofResponse,
    DocumentRelevantUnitResponse,
    DocumentTimingResponse,
    DocumentViewResponse,
)
from app.services.classification.marking_labels import get_marking_label
from app.services.classification.orchestrator import ClassificationOrchestrator
from app.services.classification.queue import extract_timings
from app.services.language_settings import get_ui_language, resolve_text
from app.services.storage.file_manager import FileManager


@dataclass(slots=True)
class RelevantUnit:
    page_number: int
    section_label: str | None
    content_type: str | None
    relevant: bool
    technical_score: int
    confidence: float
    summary: str
    citations: list[dict[str, Any]]
    evidence: list[dict[str, Any]]


@dataclass(slots=True)
class ArtifactBuildResult:
    revision: int
    original_snapshot_path: str
    marked_revision_path: str
    current_path: str
    original_sha256: str
    marked_sha256: str
    media_type: str
    duration_ms: int
    relevant_units: list[RelevantUnit]


def _sha256(path: str | Path) -> str:
    digest = hashlib.sha256()
    with Path(path).open("rb") as handle:
        for chunk in iter(lambda: handle.read(1024 * 1024), b""):
            digest.update(chunk)
    return digest.hexdigest()


def _target_unit_numbers(relevant_units: list[RelevantUnit], total_units: int) -> list[int]:
    targets = [unit.page_number for unit in relevant_units if unit.relevant]
    if targets:
        return targets
    return [1] if total_units > 0 else []


def _marking_banner(final_marking: str, *, title: str | None = None) -> str:
    return (title or get_marking_label(final_marking)).strip()


def _build_relevant_units(
    document: Document,
    page_classifications: list[dict[str, Any]] | list[Any],
) -> list[RelevantUnit]:
    page_lookup = {page.page_number: page for page in document.pages}
    units: list[RelevantUnit] = []
    for page_payload in page_classifications:
        page_number = int(getattr(page_payload, "page_number", None) or page_payload.get("page_number", 0))
        page = page_lookup.get(page_number)
        units.append(
            RelevantUnit(
                page_number=page_number,
                section_label=getattr(page, "section_label", None),
                content_type=str(getattr(page, "content_type", None) or page_payload.get("content_type") or ""),
                relevant=bool(getattr(page_payload, "relevant", None) if hasattr(page_payload, "relevant") else page_payload.get("relevant", False)),
                technical_score=int(getattr(page_payload, "technical_score", None) if hasattr(page_payload, "technical_score") else page_payload.get("technical_score", 0) or 0),
                confidence=float(getattr(page_payload, "confidence", None) if hasattr(page_payload, "confidence") else page_payload.get("confidence", 0.0) or 0.0),
                summary=str(getattr(page_payload, "summary", None) if hasattr(page_payload, "summary") else page_payload.get("summary", "")),
                citations=list(getattr(page_payload, "citations", None) if hasattr(page_payload, "citations") else page_payload.get("citations", []) or []),
                evidence=list(getattr(page_payload, "evidence", None) if hasattr(page_payload, "evidence") else page_payload.get("evidence", []) or []),
            )
        )
    units.sort(key=lambda item: item.page_number)
    return units


def _document_ext(document: Document) -> str:
    return Path(document.original_filename).suffix.lower()


def _coerce_timing_ms(value: Any) -> int | None:
    if value is None:
        return None
    try:
        return int(round(float(value)))
    except (TypeError, ValueError):
        return None


def _render_reasoning_summary(classification: Any) -> str:
    page_payloads: list[dict[str, Any]] = []
    for page_classification in getattr(classification, "page_classifications", []) or []:
        step_payload = getattr(page_classification, "step_payload", None) or {}
        payload = dict(step_payload) if isinstance(step_payload, dict) else {}
        payload.setdefault("page_number", getattr(page_classification, "page_number", 0))
        payload.setdefault("relevant", getattr(page_classification, "relevant", False))
        payload.setdefault("content_signals", {})
        payload.setdefault("tangible_hints", [])
        page_payloads.append(payload)
    return ClassificationOrchestrator.render_reasoning_summary_v2(
        getattr(classification, "reasoning_summary", None),
        final_marking=getattr(classification, "final_marking", ""),
        final_confidence=float(getattr(classification, "final_confidence", 0.0) or 0.0),
        requires_review=bool(getattr(classification, "requires_human_review", False)),
        step_results=getattr(classification, "step_results", None) or {},
        tangible_description=getattr(classification, "tangible_description", None),
        lifecycle_phase=getattr(classification, "lifecycle_phase", None),
        page_payloads=page_payloads,
    )


def _artifact_stamp_text(final_marking: str) -> str:
    return get_marking_label(final_marking).strip()


def _mark_pdf(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None, target_pages: list[int]) -> None:
    doc = fitz.open(source_path)
    banner = _artifact_stamp_text(final_marking)
    for page_number in target_pages:
        if page_number < 1 or page_number > len(doc):
            continue
        page = doc[page_number - 1]
        font_size = 6
        box_width = min(max(fitz.get_text_length(banner, fontname="helv", fontsize=font_size) + 12, 96), page.rect.width * 0.34)
        x1 = page.rect.width - 6
        x0 = max(6, x1 - box_width)
        y1 = page.rect.height - 6
        rect = fitz.Rect(x0, y1 - 14, x1, y1)
        inner_rect = fitz.Rect(rect.x0 + 3, rect.y0 + 1.5, rect.x1 - 3, rect.y1 - 1.5)
        page.draw_rect(rect, color=(0.58, 0.58, 0.58), fill=(0.97, 0.97, 0.97), width=0.5, overlay=True)
        page.insert_textbox(
            inner_rect,
            banner,
            fontsize=font_size,
            fontname="helv",
            color=(0.18, 0.18, 0.18),
            align=1,
        )
    doc.save(output_path)
    doc.close()


def _mark_image(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None) -> None:
    banner = _artifact_stamp_text(final_marking)
    with Image.open(source_path) as image:
        if image.mode != "RGBA":
            image = image.convert("RGBA")
        working = image.copy()
        overlay = Image.new("RGBA", working.size, (0, 0, 0, 0))
        draw = ImageDraw.Draw(overlay)
        font = ImageFont.load_default()
        wrapped_banner = "\n".join(textwrap.wrap(banner, width=28)) or banner
        text_bbox = draw.multiline_textbbox((0, 0), wrapped_banner, font=font, spacing=2)
        box_width = min(working.width - 12, text_bbox[2] - text_bbox[0] + 12)
        box_height = text_bbox[3] - text_bbox[1] + 8
        x0 = max(6, working.width - box_width - 6)
        y0 = max(6, working.height - box_height - 6)
        x1 = x0 + box_width
        y1 = y0 + box_height
        draw.rounded_rectangle((x0, y0, x1, y1), radius=6, fill=(245, 245, 245, 210), outline=(120, 120, 120, 210), width=1)
        draw.multiline_text((x0 + 6, y0 + 4), wrapped_banner, fill=(38, 38, 38, 255), font=font, spacing=2)
        composed = Image.alpha_composite(working, overlay)
        composed.convert("RGB").save(output_path)


def _mark_docx(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None, target_units: list[int]) -> None:
    document = DocxDocument(source_path)
    banner = _marking_banner(final_marking)
    marker = document.add_paragraph()
    marker.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    run = marker.add_run(banner)
    run.italic = True
    run.font.size = Pt(9)
    run.font.color.rgb = RGBColor(90, 90, 90)
    document.save(output_path)


def _mark_xlsx(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None, target_units: list[int]) -> None:
    workbook = load_workbook(source_path)
    banner = _marking_banner(final_marking)
    for unit in target_units:
        if unit < 1 or unit > len(workbook.worksheets):
            continue
        sheet = workbook.worksheets[unit - 1]
        target_row = max(2, sheet.max_row + 2)
        sheet.cell(row=target_row, column=1, value=banner)
    workbook.save(output_path)


def _mark_pptx(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None, target_units: list[int]) -> None:
    presentation = Presentation(source_path)
    banner = _marking_banner(final_marking)
    for unit in target_units:
        if unit < 1 or unit > len(presentation.slides):
            continue
        slide = presentation.slides[unit - 1]
        textbox = slide.shapes.add_textbox(presentation.slide_width - 260, presentation.slide_height - 32, 232, 18)
        frame = textbox.text_frame
        frame.text = banner
        if frame.paragraphs:
            frame.paragraphs[0].alignment = 1
            if frame.paragraphs[0].runs:
                frame.paragraphs[0].runs[0].font.size = Pt(10)
    presentation.save(output_path)


def _mark_text(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None) -> None:
    banner = _artifact_stamp_text(final_marking)
    text = source_path.read_text(encoding="utf-8", errors="ignore")
    output_path.write_text(f"{text}\n\n{banner}", encoding="utf-8")


def _mark_email(source_path: Path, output_path: Path, *, final_marking: str, final_confidence: float | None) -> None:
    banner = _artifact_stamp_text(final_marking)
    with source_path.open("rb") as handle:
        message = BytesParser(policy=policy.default).parse(handle)

    for part in message.walk():
        if part.get_content_type() == "text/plain" and part.get_content_disposition() is None:
            payload = part.get_content()
            if isinstance(payload, str):
                part.set_content(f"{payload}\n\n{banner}")
                break

    with output_path.open("wb") as handle:
        handle.write(message.as_bytes(policy=policy.default))


def _mark_file(
    source_path: Path,
    output_path: Path,
    *,
    document: Document,
    final_marking: str,
    final_confidence: float | None,
    target_units: list[int],
) -> None:
    ext = _document_ext(document)
    if ext == ".pdf":
        _mark_pdf(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence, target_pages=target_units)
        return
    if ext in {".png", ".jpg", ".jpeg", ".tif", ".tiff", ".webp", ".bmp"}:
        _mark_image(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence)
        return
    if ext == ".docx":
        _mark_docx(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence, target_units=target_units)
        return
    if ext == ".xlsx":
        _mark_xlsx(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence, target_units=target_units)
        return
    if ext == ".pptx":
        _mark_pptx(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence, target_units=target_units)
        return
    if ext in {".txt", ".md", ".csv"}:
        _mark_text(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence)
        return
    if ext in {".eml", ".msg"}:
        _mark_email(source_path, output_path, final_marking=final_marking, final_confidence=final_confidence)
        return
    raise ValueError(f"Format de marquage non supporte: {ext or document.media_type}")


def build_document_view(
    *,
    document: Document,
    latest_job,
    classification=None,
) -> DocumentViewResponse:
    artifact_available = bool(document.artifact_status == ArtifactStatus.READY and document.marked_artifact_path)
    document_item = DocumentListItem(
        id=document.id,
        batch_id=document.batch_id,
        client_profile_id=document.client_profile_id,
        original_filename=document.original_filename,
        relative_path=document.relative_path,
        media_type=document.media_type,
        kind=document.kind,
        status=document.status,
        page_count=document.page_count,
        duration_ms=latest_job.duration_ms if latest_job else None,
        context_hint=document.context_hint,
        processing_error=document.processing_error,
        artifact_status=document.artifact_status,
        artifact_available=artifact_available,
        latest_job_id=latest_job.id if latest_job else None,
        latest_job_status=latest_job.status if latest_job else None,
        latest_job_error=latest_job.error_message if latest_job else None,
        final_marking=classification.final_marking if classification else None,
        final_confidence=classification.final_confidence if classification else None,
        experimental=classification.experimental if classification else None,
        requires_human_review=classification.requires_human_review if classification else None,
        review_status=classification.review.status if classification and classification.review else None,
        review_priority=classification.review_priority if classification else None,
        progress=latest_job and latest_job.gate_report and latest_job.gate_report.get("progress") or None,
        created_at=document.created_at,
        updated_at=document.updated_at,
    )

    if not latest_job:
        ui_language = get_ui_language()
        return DocumentViewResponse(
            state="pending",
            message=resolve_text(
                english="Document has no associated job.",
                french="Document sans job associe.",
                language=ui_language,
            ),
            document=document_item,
            audit=DocumentAuditResponse(
                document_id=document.id,
                classification_id=classification.id if classification else None,
                job_id=None,
                latest_job_status=None,
                latest_job_error=None,
                review_status=classification.review.status if classification and classification.review else None,
                reviewer_name=classification.review.reviewer_name if classification and classification.review else None,
                reviewer_note=classification.review.reviewer_note if classification and classification.review else None,
                pipeline_version=classification.pipeline_version if classification else None,
                processing_error=document.processing_error,
                artifact_status=document.artifact_status,
                artifact_error=document.artifact_error,
            ),
            artifacts=DocumentArtifactResponse(
                status=document.artifact_status,
                version=document.artifact_version,
                has_original_snapshot=bool(document.original_snapshot_path),
                has_marked_revision=bool(document.marked_artifact_path),
                download_url=None,
                batch_download_url=None,
                relative_path=document.relative_path,
            ),
        )

    if not classification:
        state = "failed" if latest_job.status == JobStatus.FAILED or document.status == DocumentStatus.FAILED or document.processing_error else "pending"
        message = document.processing_error or latest_job.error_message
        ui_language = get_ui_language()
        return DocumentViewResponse(
            state=state,
            message=message or (
                resolve_text(
                    english="Analysis in progress",
                    french="Analyse en cours",
                    language=ui_language,
                )
                if state == "pending"
                else resolve_text(
                    english="Analysis unavailable",
                    french="Analyse indisponible",
                    language=ui_language,
                )
            ),
            document=document_item,
            audit=DocumentAuditResponse(
                document_id=document.id,
                classification_id=None,
                job_id=latest_job.id,
                latest_job_status=latest_job.status,
                latest_job_error=latest_job.error_message,
                review_status=None,
                reviewer_name=None,
                reviewer_note=None,
                pipeline_version=None,
                processing_error=document.processing_error,
                artifact_status=document.artifact_status,
                artifact_error=document.artifact_error,
            ),
            timings=DocumentTimingResponse(total_ms=latest_job.duration_ms),
            artifacts=DocumentArtifactResponse(
                status=document.artifact_status,
                version=document.artifact_version,
                has_original_snapshot=bool(document.original_snapshot_path),
                has_marked_revision=bool(document.marked_artifact_path),
                download_url=None,
                batch_download_url=None,
                relative_path=document.relative_path,
            ),
        )

    step_timing = classification.step_results.get("step_timing") if isinstance(classification.step_results, dict) else {}
    if not isinstance(step_timing, dict):
        step_timing = {}
    gate_timings = extract_timings(latest_job.gate_report) or {}
    page_scan_ms = _coerce_timing_ms(gate_timings.get("page_scan_ms"))
    if page_scan_ms is None:
        page_scan_ms = int(round(float(step_timing.get("step_a") or 0) * 1000)) or None
    decision_ms = _coerce_timing_ms(gate_timings.get("decision_ms"))
    if decision_ms is None:
        decision_seconds = sum(
            float(value)
            for key, value in step_timing.items()
            if key != "step_a" and isinstance(value, (int, float))
        )
        decision_ms = int(round(decision_seconds * 1000)) or None
    total_ms = latest_job.duration_ms
    analysis_ms = _coerce_timing_ms(gate_timings.get("analysis_ms"))
    if analysis_ms is None:
        analysis_ms = int(round(((page_scan_ms or 0) + (decision_ms or 0)))) or None

    relevant_units = _build_relevant_units(document, classification.page_classifications)
    decision = DocumentDecisionResponse(
        final_marking=classification.final_marking,
        final_confidence=classification.final_confidence,
        requires_human_review=classification.requires_human_review,
        review_priority=classification.review_priority,
        lane=classification.lane,
        jurisdiction=classification.jurisdiction,
        pipeline_version=classification.pipeline_version,
        lifecycle_phase=classification.lifecycle_phase,
        tangible_description=classification.tangible_description,
        tangible_category=classification.tangible_category,
    )
    proofs = DocumentProofResponse(
        reasoning_summary=_render_reasoning_summary(classification),
        candidate_markings=classification.candidate_markings,
        citations=classification.citations,
        evidence_items=classification.evidence_items,
    )
    audit = DocumentAuditResponse(
        document_id=document.id,
        classification_id=classification.id,
        job_id=classification.job_id,
        latest_job_status=latest_job.status,
        latest_job_error=latest_job.error_message,
        review_status=classification.review.status if classification.review else None,
        reviewer_name=classification.review.reviewer_name if classification.review else None,
        reviewer_note=classification.review.reviewer_note if classification.review else None,
        pipeline_version=classification.pipeline_version,
        processing_error=document.processing_error,
        artifact_status=document.artifact_status,
        artifact_error=document.artifact_error,
    )
    timings = DocumentTimingResponse(
        queue_ms=_coerce_timing_ms(gate_timings.get("queue_wait_ms")),
        ingestion_ms=_coerce_timing_ms(gate_timings.get("ingestion_ms")),
        page_scan_ms=page_scan_ms,
        decision_ms=decision_ms,
        analysis_ms=analysis_ms,
        marking_ms=_coerce_timing_ms(gate_timings.get("marking_ms")) or document.marking_duration_ms,
        total_ms=total_ms,
    )
    artifacts = DocumentArtifactResponse(
        status=document.artifact_status,
        version=document.artifact_version,
        has_original_snapshot=bool(document.original_snapshot_path),
        has_marked_revision=bool(document.marked_artifact_path),
        download_url=None,
        batch_download_url=None,
        relative_path=document.relative_path,
    )
    return DocumentViewResponse(
        state="ready",
        message=None,
        document=document_item,
        decision=decision,
        proofs=proofs,
        audit=audit,
        timings=timings,
        artifacts=artifacts,
        relevant_units=[
            DocumentRelevantUnitResponse(
                page_number=unit.page_number,
                section_label=unit.section_label,
                content_type=unit.content_type,
                relevant=unit.relevant,
                technical_score=unit.technical_score,
                confidence=unit.confidence,
                summary=unit.summary,
                citations=unit.citations,
                evidence=unit.evidence,
            )
            for unit in relevant_units
        ],
    )


async def build_and_persist_document_artifact(
    *,
    document: Document,
    classification,
    file_manager: FileManager,
    page_classifications: list[dict[str, Any]] | list[Any],
) -> ArtifactBuildResult:
    started = time.perf_counter()
    target_units = _build_relevant_units(document, page_classifications)
    target_pages = _target_unit_numbers(target_units, total_units=len(document.pages))
    if not target_pages:
        target_pages = [1] if document.page_count else []

    source_path = Path(document.source_path)
    if not source_path.exists():
        raise FileNotFoundError(f"Source document missing: {source_path}")

    original_snapshot_path = file_manager.snapshot_path(document.id, document.original_filename)
    original_snapshot_path.parent.mkdir(parents=True, exist_ok=True)
    shutil.copy2(source_path, original_snapshot_path)

    revision = max(1, int(document.artifact_version or 0) + 1)
    marked_revision_path = file_manager.artifact_revision_path(document.id, revision, document.original_filename)
    marked_revision_path.parent.mkdir(parents=True, exist_ok=True)
    _mark_file(
        source_path,
        marked_revision_path,
        document=document,
        final_marking=classification.final_marking,
        final_confidence=classification.final_confidence,
        target_units=target_pages,
    )

    marked_sha256 = _sha256(marked_revision_path)
    original_sha256 = _sha256(original_snapshot_path)
    duration_ms = int((time.perf_counter() - started) * 1000)

    return ArtifactBuildResult(
        revision=revision,
        original_snapshot_path=str(original_snapshot_path),
        marked_revision_path=str(marked_revision_path),
        current_path=str(marked_revision_path),
        original_sha256=original_sha256,
        marked_sha256=marked_sha256,
        media_type=document.media_type,
        duration_ms=duration_ms,
        relevant_units=target_units,
    )


def persist_document_artifact_revision(
    *,
    session,
    document: Document,
    artifact_result: ArtifactBuildResult,
) -> None:
    document.original_snapshot_path = artifact_result.original_snapshot_path
    document.marked_artifact_path = artifact_result.current_path
    document.artifact_status = ArtifactStatus.READY
    document.artifact_version = artifact_result.revision
    document.marking_duration_ms = artifact_result.duration_ms
    document.artifact_error = None

    session.add(
        DocumentArtifactRevision(
            document_id=document.id,
            revision=artifact_result.revision,
            kind=ArtifactRevisionKind.ORIGINAL_SNAPSHOT,
            filename=document.original_filename,
            media_type=document.media_type,
            file_path=artifact_result.original_snapshot_path,
            sha256=artifact_result.original_sha256,
            relative_path=document.relative_path,
            is_current=False,
        )
    )
    session.add(
        DocumentArtifactRevision(
            document_id=document.id,
            revision=artifact_result.revision,
            kind=ArtifactRevisionKind.MARKED_REVISION,
            filename=document.original_filename,
            media_type=document.media_type,
            file_path=artifact_result.marked_revision_path,
            sha256=artifact_result.marked_sha256,
            relative_path=document.relative_path,
            is_current=True,
        )
    )
