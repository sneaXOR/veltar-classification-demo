from __future__ import annotations

from pathlib import Path

import openpyxl
from docx import Document as DocxDocument
from pptx import Presentation

from app.models.enums import DocumentKind, PageContentType
from app.services.ingestion.types import IngestedDocument, IngestedPage


def _extract_docx(path: str) -> list[IngestedPage]:
    document = DocxDocument(path)
    paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
    sections = []
    chunk: list[str] = []
    page_number = 1
    for paragraph in paragraphs:
        chunk.append(paragraph)
        if len("\n".join(chunk)) > 3000:
            sections.append(
                IngestedPage(
                    page_number=page_number,
                    content_type=PageContentType.TEXT,
                    section_label=f"Section {page_number}",
                    extracted_text="\n".join(chunk),
                    metadata={"source": "docx"},
                )
            )
            chunk = []
            page_number += 1
    if chunk:
        sections.append(
            IngestedPage(
                page_number=page_number,
                content_type=PageContentType.TEXT,
                section_label=f"Section {page_number}",
                extracted_text="\n".join(chunk),
                metadata={"source": "docx"},
            )
        )
    return sections or [
        IngestedPage(
            page_number=1,
            content_type=PageContentType.TEXT,
            section_label="Section 1",
            extracted_text="",
            metadata={"source": "docx"},
        )
    ]


def _extract_xlsx(path: str) -> list[IngestedPage]:
    workbook = openpyxl.load_workbook(path, read_only=True, data_only=True)
    pages: list[IngestedPage] = []
    for index, sheet in enumerate(workbook.worksheets, start=1):
        rows: list[str] = []
        for row in sheet.iter_rows(values_only=True):
            values = [str(value).strip() for value in row if value is not None and str(value).strip()]
            if values:
                rows.append(" | ".join(values))
        pages.append(
            IngestedPage(
                page_number=index,
                content_type=PageContentType.TEXT,
                section_label=f"Sheet {sheet.title}",
                extracted_text="\n".join(rows),
                metadata={"source": "xlsx"},
            )
        )
    return pages


def _extract_pptx(path: str) -> list[IngestedPage]:
    presentation = Presentation(path)
    pages: list[IngestedPage] = []
    for index, slide in enumerate(presentation.slides, start=1):
        texts = [shape.text.strip() for shape in slide.shapes if hasattr(shape, "text") and shape.text.strip()]
        pages.append(
            IngestedPage(
                page_number=index,
                content_type=PageContentType.TEXT,
                section_label=f"Slide {index}",
                extracted_text="\n".join(texts),
                metadata={"source": "pptx"},
            )
        )
    return pages


def extract_office(path: str) -> IngestedDocument:
    suffix = Path(path).suffix.lower()
    if suffix == ".docx":
        pages = _extract_docx(path)
    elif suffix in {".xlsx", ".xlsm"}:
        pages = _extract_xlsx(path)
    elif suffix == ".pptx":
        pages = _extract_pptx(path)
    else:
        pages = [
            IngestedPage(
                page_number=1,
                content_type=PageContentType.TEXT,
                section_label="Section 1",
                extracted_text="Unsupported office format.",
                metadata={"source": "office"},
            )
        ]
    return IngestedDocument(kind=DocumentKind.OFFICE, pages=pages, metadata={"source": suffix})
